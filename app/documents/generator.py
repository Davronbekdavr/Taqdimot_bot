from datetime import datetime, timedelta, timezone
import os
import re
import shutil
import subprocess
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape


class DocumentError(RuntimeError):
    pass


def safe_filename(title: str) -> str:
    safe = re.sub(r"[^\w -]", "", title, flags=re.UNICODE).strip().replace(" ", "_")[:64]
    return safe or "AI_Studio"


class DocumentGenerator:
    def __init__(self, output_dir: str = "data/generated"):
        self.output_dir = Path(output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _id_path(self, extension: str) -> Path:
        return self.output_dir / f"{uuid4().hex}.{extension}"

    def report_docx(self, material: dict, *, body_font_pt: int, cover_details: str = "", language: str = "uz") -> Path:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Cm, Pt
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError as exc:
            raise DocumentError("python-docx_not_installed") from exc
        doc = Document()
        for section in doc.sections:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin = section.bottom_margin = Cm(2.5)
            section.left_margin = section.right_margin = Cm(3)
        for style in doc.styles:
            try:
                style.font.name = "Times New Roman"
                rfonts=style._element.get_or_add_rPr().get_or_add_rFonts()
                rfonts.set(qn("w:eastAsia"), "Times New Roman")
                rfonts.set(qn("w:cs"), "Times New Roman")
            except (AttributeError, TypeError):
                pass
        normal = doc.styles["Normal"]
        normal.font.size = Pt(body_font_pt)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.first_line_indent = Cm(1.25)
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, points in (("Title",18),("Heading 1",16),("Heading 2",14)):
            style=doc.styles[name]
            style.font.name="Times New Roman"; style.font.size=Pt(points); style.font.bold=True
            rfonts=style._element.get_or_add_rPr().get_or_add_rFonts()
            rfonts.set(qn("w:eastAsia"),"Times New Roman"); rfonts.set(qn("w:cs"),"Times New Roman")
            style.paragraph_format.keep_with_next=True

        labels={"uz":("Referat","Mundarija","Kirish","Xulosa","Foydalanilgan adabiyotlar"),"en":("Academic paper","Contents","Introduction","Conclusion","References"),"ru":("Реферат","Содержание","Введение","Заключение","Список литературы"),"tr":("Akademik çalışma","İçindekiler","Giriş","Sonuç","Kaynakça")}
        paper_label,toc_label,intro_label,conclusion_label,references_label=labels.get(language,labels["uz"])
        p=doc.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for line in cover_details.splitlines() if cover_details else []:
            if line.strip():
                q=doc.add_paragraph(line.strip()); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for _ in range(3): doc.add_paragraph()
        title=doc.add_paragraph(style="Title"); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
        title.add_run(material["title"])
        q=doc.add_paragraph(paper_label); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        doc.add_heading(toc_label,level=1)
        p=doc.add_paragraph()
        fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),'TOC \\o "1-2" \\h \\z \\u'); p._p.append(fld)
        doc.add_page_break()
        doc.add_heading(intro_label,level=1); doc.add_paragraph(material["introduction"])
        for chapter in material["chapters"]:
            doc.add_heading(chapter["heading"],level=1)
            for paragraph in re.split(r"\n\s*\n", chapter["body"]):
                if paragraph.strip(): doc.add_paragraph(paragraph.strip())
        doc.add_heading(conclusion_label,level=1); doc.add_paragraph(material["conclusion"])
        if material.get("bibliography"):
            doc.add_heading(references_label,level=1)
            for index, source in enumerate(material["bibliography"],1):
                p=doc.add_paragraph(style="Normal"); p.paragraph_format.first_line_indent=Cm(-0.5); p.paragraph_format.left_indent=Cm(0.5)
                p.add_run(f"{index}. {source}")
        for section in doc.sections:
            footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=footer.add_run(); run.font.name="Times New Roman"; run.font.size=Pt(10)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"Times New Roman")
            begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
            instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=" PAGE "
            end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
            run._r.extend([begin,instr,end])
        settings=doc.settings.element
        update=OxmlElement("w:updateFields"); update.set(qn("w:val"),"true"); settings.append(update)
        path=self._id_path("docx"); doc.save(path); return path

    def presentation_pptx(self, material: dict, design: str = "modern") -> Path:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE
            from pptx.dml.color import RGBColor
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise DocumentError("python-pptx_not_installed") from exc
        palette={"modern":((14,31,55),(0,173,181),(245,248,252)),"academic":((30,45,74),(179,139,82),(250,247,240)),"minimal":((32,40,51),(95,121,148),(255,255,255))}
        navy,accent,background=palette.get(design,palette["modern"])
        deck=Presentation(); deck.slide_width=Inches(13.333); deck.slide_height=Inches(7.5)
        blank=deck.slide_layouts[6]
        slides=material["slides"]
        for index, item in enumerate(slides):
            slide=deck.slides.add_slide(blank)
            fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=RGBColor(* (navy if index==0 else background))
            band=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(.22), Inches(7.5)); band.fill.solid(); band.fill.fore_color.rgb=RGBColor(*accent); band.line.fill.background()
            title=slide.shapes.add_textbox(Inches(.8), Inches(.65), Inches(11.7), Inches(1.1)).text_frame
            title.word_wrap=True; run=title.paragraphs[0].add_run(); run.text=item["title"]
            run.font.name="Aptos Display"; run.font.size=Pt(30 if index==0 else 26); run.font.bold=True; run.font.color.rgb=RGBColor(*(background if index==0 else navy))
            if index==0:
                deck_title=slide.shapes.add_textbox(Inches(.85), Inches(.35), Inches(10.8), Inches(.25)).text_frame.paragraphs[0]
                deck_title.text=material["title"]; deck_title.font.name="Aptos"; deck_title.font.size=Pt(10); deck_title.font.color.rgb=RGBColor(*accent)
            bullets=item.get("bullets",[])
            box=slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.3)).text_frame
            box.word_wrap=True
            for i, text in enumerate(bullets[:7]):
                p=box.paragraphs[0] if i==0 else box.add_paragraph(); p.text=text; p.level=0; p.space_after=Pt(14)
                p.font.name="Aptos"; p.font.size=Pt(22 if len(text)<90 else 18); p.font.color.rgb=RGBColor(*(background if index==0 else navy))
            footer=slide.shapes.add_textbox(Inches(10.9), Inches(7.05), Inches(1.8), Inches(.25)).text_frame.paragraphs[0]
            footer.text=f"{index+1:02d} / {len(slides):02d}"; footer.font.size=Pt(10); footer.font.color.rgb=RGBColor(*accent)
            if item.get("speaker_notes"):
                try: slide.notes_slide.notes_text_frame.text=item["speaker_notes"]
                except (AttributeError,NotImplementedError): pass
        path=self._id_path("pptx"); deck.save(path); return path

    def report_pdf(self, material: dict, *, body_font_pt: int, cover_details: str = "", language: str = "uz") -> tuple[Path,int,int]:
        font_path=os.getenv("TIMES_NEW_ROMAN_REGULAR", "")
        if not font_path or not Path(font_path).is_file():
            raise DocumentError("times_new_roman_font_missing")
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, PageBreak, TableOfContents, Flowable
        except ImportError as exc:
            raise DocumentError("reportlab_not_installed") from exc
        bold_path=os.getenv("TIMES_NEW_ROMAN_BOLD",str(Path(font_path).with_name("timesbd.ttf")))
        if not Path(bold_path).is_file():
            raise DocumentError("times_new_roman_bold_font_missing")
        pdfmetrics.registerFont(TTFont("TimesNewRoman",font_path))
        pdfmetrics.registerFont(TTFont("TimesNewRoman-Bold",bold_path))
        labels={"uz":("Referat","Mundarija","Kirish","Xulosa","Foydalanilgan adabiyotlar"),"en":("Academic paper","Contents","Introduction","Conclusion","References"),"ru":("Реферат","Содержание","Введение","Заключение","Список литературы"),"tr":("Akademik çalışma","İçindekiler","Giriş","Sonuç","Kaynakça")}
        paper_label,toc_label,intro_label,conclusion_label,references_label=labels.get(language,labels["uz"])
        path=self._id_path("pdf"); pages=[]
        styles=getSampleStyleSheet()
        title_style=ParagraphStyle("TNRTitle",parent=styles["Title"],fontName="TimesNewRoman-Bold",fontSize=18,leading=24,alignment=TA_CENTER,textColor=colors.HexColor("#172B4D"),spaceAfter=18)
        h_style=ParagraphStyle("TNRHeading",parent=styles["Heading1"],fontName="TimesNewRoman-Bold",fontSize=15,leading=19,textColor=colors.HexColor("#146C78"),spaceBefore=12,spaceAfter=8,keepWithNext=True)
        body_style=ParagraphStyle("TNRBody",parent=styles["BodyText"],fontName="TimesNewRoman",fontSize=body_font_pt,leading=body_font_pt*1.5,alignment=TA_JUSTIFY,firstLineIndent=1.25*cm,spaceAfter=7)
        class BodyEnd(Flowable):
            def __init__(self): super().__init__(); self.width=0; self.height=0
            def draw(self): pass
        story=[]
        for line in cover_details.splitlines() if cover_details else []:
            if line.strip(): story.extend([Paragraph(escape(line),ParagraphStyle("cover",parent=body_style,alignment=TA_CENTER,firstLineIndent=0)),Spacer(1,5)])
        story.extend([Spacer(1,3*cm),Paragraph(escape(material["title"]),title_style),Paragraph(escape(paper_label),body_style),PageBreak(),Paragraph(escape(toc_label),h_style)])
        toc=TableOfContents()
        toc.levelStyles=[ParagraphStyle("TNRTOC",fontName="TimesNewRoman",fontSize=body_font_pt,leading=body_font_pt*1.35,leftIndent=12,firstLineIndent=-12,textColor=colors.HexColor("#172B4D"))]
        story.append(toc)
        story.append(PageBreak()); story.extend([Paragraph(escape(intro_label),h_style),Paragraph(escape(material["introduction"]),body_style)])
        for c in material["chapters"]:
            story.append(Paragraph(escape(c["heading"]),h_style))
            story.extend(Paragraph(escape(p),body_style) for p in re.split(r"\n\s*\n",c["body"]) if p.strip())
        story.extend([Paragraph(escape(conclusion_label),h_style),Paragraph(escape(material["conclusion"]),body_style)])
        story.append(BodyEnd())
        if material.get("bibliography"):
            story.append(Paragraph(escape(references_label),h_style))
            story.extend(Paragraph(f"{i}. {escape(item)}",body_style) for i,item in enumerate(material["bibliography"],1))
        else:
            story.append(Paragraph(escape(references_label),h_style))
            story.append(Paragraph({"uz":"Tekshirilgan manbalar mavjud emas.","en":"No sources could be independently verified.","ru":"Не удалось независимо проверить источники.","tr":"Bağımsız olarak doğrulanmış kaynak bulunamadı."}[language],body_style))
        def footer(canvas,doc):
            pages.append(doc.page); canvas.saveState(); canvas.setFont("TimesNewRoman",10)
            canvas.drawCentredString(A4[0]/2,1.2*cm,str(doc.page)); canvas.restoreState()
        class ContentsDocument(BaseDocTemplate):
            def __init__(self,*args,**kwargs): super().__init__(*args,**kwargs); self.body_start=None; self.body_end=None
            def afterFlowable(self,flowable):
                if isinstance(flowable,Paragraph) and flowable.style.name=="TNRHeading":
                    title_text=flowable.getPlainText(); key=f"heading-{self.page}-{abs(hash(title_text))}"
                    self.canv.bookmarkPage(key); self.notify("TOCEntry",(0,title_text,self.page,key))
                    if title_text==intro_label and self.body_start is None: self.body_start=self.page
                elif isinstance(flowable,BodyEnd): self.body_end=self.page
        document=ContentsDocument(str(path),pagesize=A4,rightMargin=3*cm,leftMargin=3*cm,topMargin=2.5*cm,bottomMargin=2.3*cm,title=material["title"])
        frame=Frame(3*cm,2.3*cm,A4[0]-6*cm,A4[1]-4.8*cm,id="main")
        document.addPageTemplates([PageTemplate(id="paper",frames=[frame],onPage=footer)])
        document.multiBuild(story,maxPasses=4)
        total_pages=max(pages,default=0)
        body_pages=(document.body_end-document.body_start+1) if document.body_start and document.body_end else total_pages
        return path,total_pages,body_pages

    def tests_docx(self, material: dict, language: str = "uz") -> Path:
        try:
            from docx import Document
            from docx.shared import Pt
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError as exc: raise DocumentError("python-docx_not_installed") from exc
        doc=Document(); title=doc.add_heading(material["title"],0); title.alignment=1
        for style in doc.styles:
            try:
                style.font.name="Times New Roman"; style.font.size=Pt(12); style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"Times New Roman")
            except (AttributeError,TypeError): pass
        for i,item in enumerate(material["questions"],1):
            doc.add_heading(f"{i}. {item['question']}",level=2)
            for option in item["options"]: doc.add_paragraph(option,style="List Bullet")
            p=doc.add_paragraph(); r=p.add_run("Answer: "+item["answer"]+". "+item["explanation"]); r.italic=True
        path=self._id_path("docx"); doc.save(path); return path

    def presentation_pdf(self, material: dict, design: str = "modern") -> tuple[Path,int]:
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_LEFT
            from reportlab.lib.pagesizes import landscape
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        except ImportError as exc: raise DocumentError("reportlab_not_installed") from exc
        palettes={"modern":("#0E1F37","#00ADB5","#F5F8FC"),"academic":("#1E2D4A","#B38B52","#FAF7F0"),"minimal":("#202833","#5F7994","#FFFFFF")}
        dark,accent,light=palettes.get(design,palettes["modern"])
        size=landscape((13.333*inch,7.5*inch)); path=self._id_path("pdf"); count=[]
        styles=getSampleStyleSheet(); title=ParagraphStyle("slide_title",parent=styles["Title"],fontName="Helvetica-Bold",fontSize=27,leading=33,textColor=colors.HexColor(dark),spaceAfter=25)
        bullet=ParagraphStyle("slide_bullet",parent=styles["BodyText"],fontName="Helvetica",fontSize=20,leading=29,leftIndent=18,spaceAfter=14,textColor=colors.HexColor(dark))
        story=[]
        for index,slide in enumerate(material["slides"]):
            story.append(Paragraph(escape(material["title"]),ParagraphStyle("deck_label",parent=styles["Normal"],textColor=colors.HexColor(accent),fontSize=10)))
            story.extend([Spacer(1,.35*inch),Paragraph(escape(slide["title"]),title)])
            for item in slide["bullets"]: story.append(Paragraph("•  "+escape(item),bullet))
            if index<len(material["slides"])-1: story.append(PageBreak())
        def footer(canvas,doc):
            count.append(1); canvas.saveState(); canvas.setFillColor(colors.HexColor(accent)); canvas.rect(0,0,.14*inch,size[1],stroke=0,fill=1)
            canvas.setFont("Helvetica",10); canvas.drawRightString(size[0]-.45*inch,.3*inch,f"{doc.page:02d} / {len(material['slides']):02d}"); canvas.restoreState()
        SimpleDocTemplate(str(path),pagesize=size,leftMargin=.8*inch,rightMargin=.55*inch,topMargin=.65*inch,bottomMargin=.65*inch).build(story,onFirstPage=footer,onLaterPages=footer)
        return path,len(count)

    @staticmethod
    def libreoffice_pdf(source: Path) -> tuple[Path,int] | None:
        executable=shutil.which("soffice") or shutil.which("libreoffice")
        if not executable: return None
        out_dir=source.parent
        proc=subprocess.run([executable,"--headless","--convert-to","pdf","--outdir",str(out_dir),str(source)],capture_output=True,text=True,timeout=90)
        candidate=out_dir/(source.stem+".pdf")
        if proc.returncode or not candidate.exists(): return None
        try:
            from pypdf import PdfReader
            count=len(PdfReader(str(candidate)).pages)
        except ImportError:
            count=0
        return candidate,count


def file_expiry(days: int = 30) -> str:
    return (datetime.now(timezone.utc)+timedelta(days=days)).isoformat(timespec="seconds")
